"""企業文件檢索（RAG）。

上傳 → 解析文字 → 切成段落 → Gemini 向量化 → 存進 PostgreSQL + pgvector。
Agent 整理流程、AI 介入探索時，依對話內容或工作節點檢索最相關的段落，並附上文件來源。
向量服務無法使用時（未設定金鑰、額度用盡）改用關鍵字比對，流程不中斷。
"""

import io
import logging
import re

from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from . import llm
from .models import EMBED_DIM, DocChunk, Document

log = logging.getLogger("bizpilot")

KINDS = ("SOP", "FAQ", "政策", "商品資料", "其他")
MAX_BYTES = 5 * 1024 * 1024
CHUNK_SIZE = 400  # 每段約 400 字：夠涵蓋一條規定，又不會混進太多不相關內容
# 向量相似度低於此值視為不相關（示範文件實測：相關段落約 0.69 以上，無關問題約 0.61）
MIN_SCORE = 0.65
MAX_GAP = 0.1  # 與最相關段落差距太大的也不列入，避免湊數
HEADING = re.compile(r"^(#+\s|[一二三四五六七八九十]+、|第[一二三四五六七八九十\d]+[條章節])")


class DocumentError(ValueError):
    """上傳的檔案無法使用；訊息直接顯示給企業。"""


def extract_text(filename: str, raw: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ("txt", "md", "csv"):
        for enc in ("utf-8-sig", "cp950"):  # 台灣常見的 Big5 文字檔
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
        raise DocumentError("無法辨識文字編碼，請另存成 UTF-8 後再上傳。")
    if ext == "docx":
        from docx import Document as Docx

        d = Docx(io.BytesIO(raw))
        lines = [p.text for p in d.paragraphs]
        for table in d.tables:
            lines += [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        return "\n".join(lines)
    if ext == "pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages)
    raise DocumentError("目前支援 PDF、Word（.docx）與文字檔（.txt、.md、.csv）。")


def _pieces(line: str, size: int) -> list[str]:
    """過長的一行依句號切開，每塊不超過 size。"""
    if len(line) <= size:
        return [line]
    out, cur = [], ""
    for sent in re.split(r"(?<=[。！？!?；;])", line):
        while len(sent) > size:  # 沒有標點的超長句直接截斷
            out.append(sent[:size])
            sent = sent[size:]
        if len(cur) + len(sent) > size and cur:
            out.append(cur)
            cur = ""
        cur += sent
    return out + ([cur] if cur else [])


def chunk_text(text: str, size: int = CHUNK_SIZE) -> list[str]:
    """依標題與段落切段：遇到標題就另起一段，其餘逐行累積到約 size 字。
    只有標題、還沒有內文的段落不單獨成段，會跟下一個標題的內容合併。"""
    chunks, cur, has_body = [], "", False
    for raw_line in text.splitlines():
        line = re.sub(r"[ \t　]+", " ", raw_line).strip()
        if not line:
            continue
        heading = bool(HEADING.match(line))
        if heading and has_body:
            chunks.append(cur)
            cur, has_body = "", False
        has_body = has_body or not heading
        for piece in _pieces(line, size):
            if cur and len(cur) + len(piece) + 1 > size:
                chunks.append(cur)
                cur = ""
            cur = f"{cur}\n{piece}" if cur else piece
    if cur:
        chunks.append(cur)
    return chunks


def _embed(texts: list[str], task: llm.EmbedTask) -> list[list[float]] | None:
    try:
        vectors = llm.embed(texts, task, EMBED_DIM)
    except Exception as e:  # 未設定金鑰或 Gemini 忙碌：改用關鍵字檢索
        log.warning("文件向量化失敗，改用關鍵字檢索：%s", e)
        return None
    return vectors if len(vectors) == len(texts) else None


def add_document(
    db: Session, tax_id: str, filename: str, kind: str, text: str, vectors: list[list[float]] | None = None
) -> Document:
    """切段並存入資料庫；vectors 未提供時呼叫 Gemini 向量化（示範文件會帶入預存向量）。"""
    chunks = chunk_text(text)
    if not chunks:
        raise DocumentError("這份文件讀不到文字（可能是掃描成圖片的 PDF）。")
    if vectors is None or len(vectors) != len(chunks):
        vectors = _embed(chunks, "RETRIEVAL_DOCUMENT")
    doc = Document(tax_id=tax_id, filename=filename, kind=kind, chars=len(text), embedded=vectors is not None)
    db.add(doc)
    db.flush()
    for i, t in enumerate(chunks):
        db.add(DocChunk(document_id=doc.id, tax_id=tax_id, idx=i + 1, text=t, embedding=vectors[i] if vectors else None))
    return doc


def upload(db: Session, tax_id: str, filename: str, kind: str, raw: bytes) -> Document:
    if kind not in KINDS:
        raise DocumentError("文件類型不正確。")
    if len(raw) > MAX_BYTES:
        raise DocumentError("檔案超過 5 MB，請拆成較小的檔案。")
    try:
        text = extract_text(filename, raw)
    except DocumentError:
        raise
    except Exception:
        raise DocumentError("檔案無法讀取，可能已損毀或有密碼保護。")
    # 同名檔案視為更新版本，取代舊的
    db.execute(delete(Document).where(Document.tax_id == tax_id, Document.filename == filename))
    return add_document(db, tax_id, filename, kind, text)


def list_documents(db: Session, tax_id: str) -> list[dict]:
    rows = db.execute(
        select(Document, func.count(DocChunk.id))
        .outerjoin(DocChunk, DocChunk.document_id == Document.id)
        .where(Document.tax_id == tax_id)
        .group_by(Document.id)
        .order_by(Document.id)
    ).all()
    return [
        {
            "id": d.id,
            "filename": d.filename,
            "kind": d.kind,
            "chars": d.chars,
            "chunks": n,
            "embedded": d.embedded,
            "created_at": d.created_at,
        }
        for d, n in rows
    ]


def _hit(c: DocChunk, d: Document, score: float) -> dict:
    return {
        "filename": d.filename,
        "kind": d.kind,
        "idx": c.idx,
        "text": c.text,
        "score": round(score, 3),
    }


def _grams(s: str) -> set[str]:
    """中文取相鄰兩字、英數取整個詞，作為關鍵字比對單位。"""
    s = s.lower()
    words = set(re.findall(r"[a-z0-9]{2,}", s))
    for run in re.findall(r"[一-鿿]+", s):
        words |= {run[i : i + 2] for i in range(len(run) - 1)} if len(run) > 1 else {run}
    return words


def _keyword_search(db: Session, tax_id: str, query: str, k: int) -> list[dict]:
    q = _grams(query)
    if not q:
        return []
    scored = []
    for c, d in db.execute(select(DocChunk, Document).join(Document).where(DocChunk.tax_id == tax_id)).all():
        score = len(q & _grams(c.text)) / len(q)
        if score >= 0.2:
            scored.append((score, c, d))
    scored.sort(key=lambda x: -x[0])
    return [_hit(c, d, s) for s, c, d in scored[:k]]


def search(db: Session, tax_id: str, query: str, k: int = 4) -> list[dict]:
    """找出與 query 最相關的文件段落；企業沒有文件時回傳空清單、不呼叫 Gemini。"""
    query = query.strip()
    if not query or not db.scalar(select(func.count()).select_from(DocChunk).where(DocChunk.tax_id == tax_id)):
        return []
    has_vectors = db.scalar(
        select(func.count()).select_from(DocChunk).where(DocChunk.tax_id == tax_id, DocChunk.embedding.is_not(None))
    )
    qv = _embed([query], "RETRIEVAL_QUERY") if has_vectors else None
    if qv:
        dist = DocChunk.embedding.cosine_distance(qv[0])
        rows = db.execute(
            select(DocChunk, Document, dist)
            .join(Document)
            .where(DocChunk.tax_id == tax_id, DocChunk.embedding.is_not(None))
            .order_by(dist)
            .limit(k)
        ).all()
        top = 1 - rows[0][2] if rows else 0
        return [_hit(c, d, 1 - dd) for c, d, dd in rows if 1 - dd >= max(MIN_SCORE, top - MAX_GAP)]
    return _keyword_search(db, tax_id, query, k)


def as_prompt(hits: list[dict]) -> str:
    """給模型看的文件段落，附編號方便對照來源。"""
    if not hits:
        return ""
    body = "\n".join(f"[{i + 1}]《{h['filename']}》第 {h['idx']} 段：\n{h['text']}" for i, h in enumerate(hits))
    return f"企業提供的文件（依相關程度排序）：\n{body}\n"
