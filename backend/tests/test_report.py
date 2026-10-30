"""需求草案：彙整內容、待確認事項、匯出檔案。Gemini 以替身取代。"""

import io
from urllib.parse import unquote

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import text

from app import llm
from app.allocation import Exploration, ModeOption
from app.db import engine
from app.main import app
from app.sandbox.engine import CONDITIONS

client = TestClient(app)


@pytest.fixture(autouse=True)
def cleanup():
    with engine.connect() as conn:
        last = conn.execute(text("select coalesce(max(id), 0) from projects")).scalar()
    yield
    with engine.begin() as conn:
        conn.execute(text("delete from projects where id > :last"), {"last": last})


def exploration(**kw):
    def opt(mode):
        return ModeOption(
            mode=mode, ai_does=f"{mode}-AI", human_does=f"{mode}-人", required_data=["訂單編號"],
            permissions=["讀取訂單"], handoff=["查無訂單時轉客服"], risk="r", fit="可行", reason="",
        )

    return Exploration(ai_tasks=["比對訂單"], goal="縮短查詢時間", options=[opt(m) for m in ("human", "assist", "review", "auto")], recommended="assist")


@pytest.fixture
def project(monkeypatch):
    monkeypatch.setattr(llm, "generate_json", lambda *a, **k: exploration())
    client.get("/api/companies/00000000")
    pid = client.post("/api/projects/demo").json()["id"]
    client.post(f"/api/projects/{pid}/nodes/lookup/explore")
    client.put(f"/api/projects/{pid}/allocations/lookup", json={"mode": "assist"})
    return pid


def report(pid):
    return client.get(f"/api/projects/{pid}/report").json()


def test_report_uses_chosen_mode_details(project):
    r = report(project)
    lookup = next(a for a in r["allocations"] if a["step"] == "查詢訂單資料")
    assert lookup["mode"] == "AI 輔助" and lookup["ai_does"] == "assist-AI"
    assert lookup["handoff"] == ["查無訂單時轉客服"]  # 每一步自己的交給人情況，重點區拿來當風險對策
    assert "訂單編號" in r["data"]
    assert any("讀取訂單" in p for p in r["permissions"])


def test_untested_project_lists_todo(project):
    todo = report(project)["todo"]
    assert "尚未執行情境沙盒測試。" in todo
    assert any("尚未由企業確認" in t for t in todo)


def test_assumed_conditions_and_gaps_become_todo(project):
    client.post(f"/api/projects/{project}/sandbox/run", json={"config": {"docs_ready": True}})
    r = report(project)
    assert r["testing"]["tested"] and not r["testing"]["stale"]
    assert any("假設「已提供 SOP" in t for t in r["todo"])
    assert any(t.startswith("[") for t in r["todo"])  # 缺口也列入待確認


def test_changing_allocation_marks_test_stale(project):
    client.post(f"/api/projects/{project}/sandbox/run", json={"config": {k: True for k in CONDITIONS}})
    client.put(f"/api/projects/{project}/allocations/notify", json={"mode": "auto"})
    r = report(project)
    assert r["testing"]["stale"]
    assert any("重新執行情境測試" in t for t in r["todo"])


def test_markdown_download(project):
    res = client.get(f"/api/projects/{project}/report/markdown")
    assert res.status_code == 200
    assert "AI 導入需求草案" in unquote(res.headers["content-disposition"])
    body = res.content.decode("utf-8")
    for heading in ["1. 專案概要", "2. 專案摘要", "3. 人力與 AI 分配", "4. 現行工作流程", "5. AI 所需資料", "6. AI 操作權限", "7. 主要風險"]:
        assert f"## {heading}" in body
    assert body.startswith("> 智行進業 BizPilot AI") and body.rstrip().endswith("也不會實際退款或開通權限。")
    assert "查詢訂單資料" in body


def test_docx_download_opens_in_word(project):
    res = client.get(f"/api/projects/{project}/report/docx")
    doc = Document(io.BytesIO(res.content))
    texts = [p.text for p in doc.paragraphs]
    assert any("AI 導入需求草案" in t for t in texts)
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["步驟", "分工方式", "AI 負責", "人員負責"]
    section = doc.sections[0]
    assert "智行進業 BizPilot AI" in section.header.paragraphs[0].text
    assert section.footer.paragraphs[0].text.startswith("免責聲明")


def fake_brief(steps: list[str], extra: str = ""):
    """重點區短句的 AI 替身；extra 模擬 AI 多寫了草案裡沒有的步驟。"""

    def generate_json(prompt, schema, system=None):
        ai = [{"step": s, "text": f"{s}的短句"} for s in steps + ([extra] if extra else [])]
        return schema.model_validate({"ai": ai, "human": "主管負責例外。"})

    return generate_json


def test_ai_summary_is_saved_and_shown(project, monkeypatch):
    monkeypatch.setattr(llm, "generate", lambda prompt, system=None: ("這是摘要。", "fake"))
    monkeypatch.setattr(llm, "generate_json", fake_brief([]))
    r = client.post(f"/api/projects/{project}/report/summary").json()
    assert r["ai_summary"] == "這是摘要。"
    assert "AI 撰寫" in client.get(f"/api/projects/{project}/report/markdown").content.decode()


def test_ai_brief_keeps_only_ai_steps(project, monkeypatch):
    """短句只保留草案中有 AI 參與的步驟；AI 多寫的步驟會被丟掉。"""
    r = client.get(f"/api/projects/{project}/report").json()
    assert r["ai_brief"] is None  # 還沒請 AI 撰寫
    ai_steps = [a["step"] for a in r["allocations"] if a["mode"] != "人工執行"]
    monkeypatch.setattr(llm, "generate", lambda prompt, system=None: ("摘要。", "fake"))
    monkeypatch.setattr(llm, "generate_json", fake_brief(ai_steps, extra="不存在的步驟"))
    r = client.post(f"/api/projects/{project}/report/summary").json()
    assert [s["step"] for s in r["ai_brief"]["ai"]] == ai_steps
    assert r["ai_brief"]["human"] == "主管負責例外。"


@pytest.mark.parametrize("company_name, shown", [("阿明電器行", "阿明電器行"), ("", "我的公司")])
def test_project_without_tax_id(company_name, shown):
    """不查統編也能開始；公司名稱可自填，沒填就叫「我的公司」，草案不顯示統編。"""
    res = client.post("/api/projects", json={"company_name": company_name, "name": "退貨改善", "work": "處理退貨"})
    assert res.status_code == 200
    p = res.json()
    try:
        assert p["company_name"] == shown and p["tax_id"].startswith("M")
        body = client.get(f"/api/projects/{p['id']}/report/markdown").content.decode("utf-8")
        assert f"- 企業：{shown}\n" in body
    finally:
        with engine.begin() as conn:
            conn.execute(text("delete from projects where id = :id"), {"id": p["id"]})
            conn.execute(text("delete from companies where tax_id = :t"), {"t": p["tax_id"]})
